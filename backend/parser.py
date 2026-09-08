# -*- coding: utf-8 -*-
"""文档解析：提取纯文本与基础信息。支持 PDF/DOCX/DOC/TXT/MD/XLSX/PPTX。
OCR（扫描件）为可选组件：检测到 paddleocr 才启用，否则给出明确提示。"""
import os

from . import config

SUPPORTED_EXTS = config.SUPPORTED_EXTS


class ParseResult:
    def __init__(self, ok=False, text="", msg=""):
        self.ok = ok
        self.text = text
        self.msg = msg


def parse_file(path: str) -> ParseResult:
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED_EXTS:
        return ParseResult(False, "", f"不支持的文件类型: {ext}")
    try:
        if ext == ".pdf":
            return _parse_pdf(path)
        if ext in (".docx",):
            return _parse_docx(path)
        if ext == ".doc":
            return _parse_doc_legacy(path)
        if ext == ".xlsx":
            return _parse_xlsx(path)
        if ext == ".pptx":
            return _parse_pptx(path)
        if ext in (".txt", ".md"):
            return _parse_text(path)
    except Exception as e:  # 容错：损坏/加密文件不崩溃
        return ParseResult(False, "", f"解析失败: {e}")
    return ParseResult(False, "", "未知错误")


def _parse_pdf(path):
    import pdfplumber
    texts = []
    with pdfplumber.open(path) as pdf:
        # 加密 PDF 会在 open 或访问 pages 时抛错
        for page in pdf.pages:
            t = page.extract_text() or ""
            texts.append(t)
    full = "\n".join(texts).strip()
    if len(full) < 20:
        return _parse_pdf_with_ocr(path, full)
    return ParseResult(True, full, "")


def _parse_pdf_with_ocr(path, fallback_text):
    """扫描版 PDF：优先本地 OCR（可选组件），未安装则明确提示。"""
    try:
        from paddleocr import PaddleOCR  # noqa
        import fitz  # PyMuPDF 渲染页面为图片
    except ImportError:
        return ParseResult(
            False, fallback_text,
            "该 PDF 无可提取文本层（扫描件）。如需 OCR 识别，请安装可选组件: "
            "pip install paddleocr paddlepaddle pymupdf（约1GB，低配机不建议）",
        )
    try:
        ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
        doc = fitz.open(path)
        out = []
        for page in doc:
            pix = page.get_pixmap(dpi=150)
            img_path = os.path.join(config.DATA_DIR, "_ocr_tmp.png")
            pix.save(img_path)
            res = ocr.ocr(img_path, cls=True)
            for line in (res or []):
                for item in (line or []):
                    out.append(item[1][0])
            if os.path.exists(img_path):
                os.remove(img_path)
        text = "\n".join(out).strip()
        return ParseResult(True, text, "OCR 识别完成") if text else ParseResult(False, "", "OCR 未识别到内容")
    except Exception as e:
        return ParseResult(False, fallback_text, f"OCR 失败: {e}")


def _parse_docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    full = "\n".join(parts).strip()
    return ParseResult(True, full, "") if full else ParseResult(False, "", "文档内容为空")


def _parse_doc_legacy(path):
    """旧版 .doc：OLE 结构尽力提取（UTF-16LE 文本流），失败则提示另存为 .docx。"""
    try:
        import olefile
        if not olefile.isOleFile(path):
            return ParseResult(False, "", "非有效 .doc 文件")
        ole = olefile.OleFileIO(path)
        if ole.exists("WordDocument"):
            data = ole.openstream("WordDocument").read()
            text = data.decode("utf-16-le", errors="ignore")
            # 过滤控制字符，保留中文与可打印 ASCII
            filtered = "".join(ch for ch in text if ch.isprintable() or ch in "\n\t")
            filtered = filtered.strip()
            if len(filtered) > 10:
                return ParseResult(True, filtered, "旧版格式尽力提取")
        return ParseResult(False, "", ".doc 旧格式提取失败，建议另存为 .docx 后导入")
    except ImportError:
        return ParseResult(False, "", ".doc 需可选组件 olefile：pip install olefile，或另存为 .docx")
    except Exception as e:
        return ParseResult(False, "", f".doc 解析失败: {e}")


def _parse_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"【工作表 {ws.title}】")
        for row in ws.iter_rows(values_only=True):
            vals = [str(c) for c in row if c is not None and str(c).strip()]
            if vals:
                parts.append(" | ".join(vals))
    wb.close()
    full = "\n".join(parts).strip()
    return ParseResult(True, full, "") if full else ParseResult(False, "", "表格内容为空")


def _parse_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"【第{i}页】")
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
    full = "\n".join(parts).strip()
    return ParseResult(True, full, "") if full else ParseResult(False, "", "演示文稿内容为空")


def _parse_text(path):
    # 依次尝试常见编码：UTF-8 / GBK / UTF-16
    raw = open(path, "rb").read()
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            text = raw.decode(enc)
            return ParseResult(True, text.strip(), "")
        except (UnicodeDecodeError, ValueError):
            continue
    return ParseResult(False, "", "文本编码无法识别")


def make_summary(text: str, limit: int = 200) -> str:
    """取文档开头作为摘要（去掉空行后）。"""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    s = "\n".join(lines)
    return s[:limit]
